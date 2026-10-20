"""Reading a register. The file is read once; its SHA-256 is taken from those bytes, and openpyxl parses that
in-memory copy twice (formulas, then cached values). Nothing here opens the file for writing."""
import csv
import hashlib
import io
import os
import warnings

from .packs import EXTRA, SetupError

FORMATS = {".xlsx": "xlsx", ".xlsm": "xlsx", ".csv": "csv"}
REFUSED = {".xls": "an Excel 97-2003 workbook", ".xlsb": "a binary Excel workbook",
           ".ods": "an OpenDocument spreadsheet", ".numbers": "a Numbers file"}
OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # encrypted workbooks and old .xls files start with this


class Sheet(object):
    def __init__(self, name, values, formulas):
        self.name = name
        self.values = values  # rows of cell values; a formula cell holds its cached value
        self.formulas = formulas  # zero-based (row, column) of every formula cell; None for a CSV


class Register(object):
    def __init__(self, path, fmt, sha256, sheets, note=None):
        self.file = os.path.basename(path)
        self.format = fmt
        self.sha256 = sha256
        self.sheets = sheets
        self.note = note


def read(path):
    name = os.path.basename(path)
    ext = os.path.splitext(path)[1].lower()
    if ext in REFUSED:
        raise SetupError("%s is %s. Save it as .xlsx or .csv, then lint that copy." % (name, REFUSED[ext]))
    if ext not in FORMATS:
        raise SetupError("%s: goslint reads .xlsx, .xlsm and .csv files." % name)
    try:
        with open(path, "rb") as f:
            data = f.read()
    except OSError as e:
        raise SetupError("cannot read %s: %s" % (path, e.strerror or e))
    sha = hashlib.sha256(data).hexdigest()
    if FORMATS[ext] == "csv":
        sheets, note = _csv(data, name)
        return Register(path, "csv", sha, sheets, note)
    return Register(path, ext[1:], sha, _xlsx(data, name))


def _xlsx(data, name):
    try:
        import openpyxl
    except ImportError:
        raise SetupError(EXTRA)
    if data[:8] == OLE:
        raise SetupError("%s is password-protected or is an old .xls file with a new name; goslint cannot open it." % name)
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")  # openpyxl warns about features it would drop on save; goslint never saves
            wf = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
            wv = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    except Exception as e:
        raise SetupError("cannot read %s as a workbook (%s)." % (name, e.__class__.__name__))
    sheets = []
    for ws in wf.worksheets:
        formulas = set()
        for r, row in enumerate(ws.iter_rows()):
            for c, cell in enumerate(row):
                if cell.data_type == "f":
                    formulas.add((r, c))
        values = [list(row) for row in wv[ws.title].iter_rows(values_only=True)]
        sheets.append(Sheet(ws.title, values, formulas))
    return sheets


def _csv(data, name):
    note = None
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
        note = "read as Windows-1252 text, not UTF-8"
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = csv.reader(io.StringIO(text, newline=""), dialect)
    values = [[v if v.strip() else None for v in row] for row in rows]
    return [Sheet(os.path.splitext(name)[0], values, None)], note
