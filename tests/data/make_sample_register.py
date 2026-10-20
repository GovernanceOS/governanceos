"""Build the sample register goslint's tests use: nebula-mutual-register.xlsx and its CSV export.

    python tests/data/make_sample_register.py

Nebula Mutual is a fictional institution constructed for illustration; any resemblance to a real
organization is unintended. Its systems come from GovernanceOS's Nebula Mutual demo records (the
public sample evidence pack holds six of the seven); people are roles, never names. The register is the kind a practitioner keeps by hand, with the gaps such
registers have: two E-23 Appendix 1 columns missing, empty cells, a placeholder, a past next review
date, an ambiguous date and two next review dates derived by formula.

openpyxl writes formulas without their results, so the formula cells get the cached values Excel
would have saved; the save time openpyxl stamps and every zip entry's timestamp are fixed, so a
rebuild with the same openpyxl gives the same bytes.
Needs openpyxl.
"""
import csv
import datetime
import io
import os
import re
import zipfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "nebula-mutual-register.xlsx")
CSV = os.path.join(HERE, "nebula-mutual-register.csv")
STAMP = (2026, 10, 19, 0, 0, 0)
LABEL = ("Nebula Mutual model inventory. Nebula Mutual is a fictional institution constructed for illustration; "
         "any resemblance to a real organization is unintended.")
HEADERS = ["Model ID", "Model name", "Description", "Risk rating", "Model owner", "Developer", "Origin", "Version",
           "Deployment date", "Validator", "Approved by", "Data sources", "Approved uses", "Last review date",
           "Monitoring status", "Next review date"]
D = datetime.date
NEXT_FROM_LAST = "=N{r}+365"  # the practitioner's own formula, read as derived; goslint computes nothing

ROWS = [
    ["nebula-mutual/customer-assistant", "Customer assistant",
     "Agent on an LLM from Meridian Language Systems, with policy-lookup and claim-status tools", "Tier 1", "Actuary",
     "Meridian Language Systems", "Vendor", "2.1", D(2026, 9, 2), "Validator", "CRO",
     "Policy and claim records, read through lookups", "Answering policyholder questions on policies and claims",
     D(2026, 6, 11), "Monitored, no open exceptions", D(2027, 3, 2)],
    ["nebula-mutual/claims-routing", "Claims routing", "Classification model routing new claims to handling queues",
     "Tier 2", "Actuary", "In house", "Internal", "1.4", D(2026, 8, 30), "Validator", "CRO", "Claim intake records",
     "Routing new claims to handling queues", D(2026, 7, 28), "Monitored, no open exceptions", NEXT_FROM_LAST],
    ["nebula-mutual/underwriting-triage", "Underwriting triage",
     "Vendor model ordering new applications for underwriter review", "Tier 2", "Actuary", "Corvane", "Vendor", "3.0",
     None, "Validator", None, "Application data", "Ordering new applications for underwriter review",
     D(2026, 8, 21), "Not in production", NEXT_FROM_LAST],
    ["nebula-mutual/document-ocr", "Document OCR", "Vendor service reading fields from claim documents", "Tier 1",
     "Actuary", "Northlight", "Vendor", "5.2", D(2025, 11, 11), "Validator", "CRO", "Scanned claim documents",
     "Reading fields from claim documents", D(2025, 9, 15), "Monitored, 1 open exception", D(2026, 9, 15)],
    ["nebula-mutual/pricing-model-v3", "Pricing model v3", "Pricing model for individual life products", "Tier 3",
     "Actuary", "In house", "Internal", "3.0", None, "Validator", None, "Policy and experience data",
     "Pricing new individual life policies", D(2026, 6, 5), "Not in production", "03/04/2027"],
    ["nebula-mutual/lapse-prediction", "Policy lapse prediction", "Model predicting policy lapse, submitted for intake",
     None, "Actuary", "TBD", "Internal", "0.1", None, None, None, "Policy records", None, None, "Not in production",
     None],
    ["nebula-mutual/fraud-flags", "Fraud flags", "Rules and model flagging claims for review, retired 2026-05-20",
     "Tier 2", "Actuary", "In house", "Internal", "2.3", D(2025, 2, 24), "Validator", "CRO", "Claim records",
     "Flagging claims for review", D(2025, 1, 17), "Retired 2026-05-20", "n/a, retired"],
]
FIRST = 4  # the first data row; row 3 holds the headers


def serial(d):
    return (d - D(1899, 12, 30)).days


def values():
    """Each row as Excel would show it: formulas replaced by their results."""
    out = []
    for i, row in enumerate(ROWS):
        r = FIRST + i
        shown = list(row)
        if row[15] == NEXT_FROM_LAST:
            shown[15] = row[13] + datetime.timedelta(days=365)
        out.append((r, row, shown))
    return out


def build_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model inventory"
    ws["A1"] = LABEL
    ws["A2"] = "Register owner: Data lead"
    ws.append(HEADERS)
    cached = {}
    for r, row, shown in values():
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c)
            if v == NEXT_FROM_LAST:
                cell.value = v.format(r=r)
                cached[cell.coordinate] = serial(shown[c - 1])
            else:
                cell.value = v
            if isinstance(shown[c - 1], D):
                cell.number_format = "yyyy-mm-dd"
    lists = wb.create_sheet("Lists")
    lists.append(["Risk tiers", "Origin"])
    for t, o in (("Tier 1", "Internal"), ("Tier 2", "Vendor"), ("Tier 3", None)):
        lists.append([t, o])
    wb.properties.creator = "GovernanceOS sample"
    wb.properties.created = wb.properties.modified = datetime.datetime(2026, 10, 16, 16, 0, 0)
    raw = io.BytesIO()
    wb.save(raw)

    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                for coord, v in cached.items():
                    text, n = re.subn(r'(<c r="%s"[^>]*><f>[^<]*</f>)<v\s*/?>(</v>)?' % coord, r"\g<1><v>%d</v>" % v, text)
                    assert n == 1, coord
                data = text.encode("utf-8")
            if item.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-10-16T16:00:00Z", data)
            info = zipfile.ZipInfo(item.filename, STAMP)
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, data)
    with open(XLSX, "wb") as f:
        f.write(out.getvalue())


def build_csv():
    """The register as Excel's "CSV UTF-8" export would write it: values only, dates as displayed."""
    with open(CSV, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow([LABEL])
        w.writerow(["Register owner: Data lead"])
        w.writerow(HEADERS)
        for _, _, shown in values():
            w.writerow(["" if v is None else v.isoformat() if isinstance(v, D) else v for v in shown])


if __name__ == "__main__":
    build_xlsx()
    build_csv()
    print("wrote %s and %s" % (XLSX, CSV))
